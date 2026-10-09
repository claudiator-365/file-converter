import json, sys, openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
m = json.load(open(sys.argv[1])); wb = openpyxl.Workbook(); wb.remove(wb.active)
for s in m:
    ws = wb.create_sheet(s['name'])
    for c in s['cells']:
        x = ws.cell(c['r'], c['c'])
        if 'value' in c: x.value = c['value']
        if 'numFmt' in c: x.number_format = c['numFmt']
        f = c.get('font', {}); x.font = Font(name=f.get('name'), size=f.get('size'))
        a = c.get('alignment', {}); x.alignment = Alignment(horizontal=a.get('horizontal'), vertical=a.get('vertical'), wrap_text=a.get('wrapText'))
        if 'border' in c: t = Side(style='thin'); x.border = Border(top=t, left=t, bottom=t, right=t)
    for a, b, c2, d in s['merges']: ws.merge_cells(start_row=a, start_column=b, end_row=c2, end_column=d)
    for i, v in s['cols'].items(): ws.column_dimensions[L(int(i))].width = v['width']
    for i, v in s['rows'].items(): ws.row_dimensions[int(i)].height = v['height']
wb.save(sys.argv[2]); print('wrote', sys.argv[2])
